#!/usr/bin/env python3
import argparse
import json
import os
import re
import sys
from pathlib import Path
from typing import List, Dict, Any, Tuple

FIGURE_MARK_RE = re.compile(r"(?i)\b(?:fig(?:ure)?\.?\s*\d+|图\s*\d+)\b")
MARKDOWN_IMAGE_RE = re.compile(r"!\[[^\]]*\]\(([^)]+)\)")
IMAGE_PATH_RE = re.compile(r"(?i)([\w./-]+\.(?:png|jpg|jpeg|svg|webp|gif))")

LAYOUT_RULES: List[Tuple[re.Pattern[str], str, str]] = [
    (re.compile(r"(?i)对比|消融|ablation|baseline"), "pin-3", "命中对比/消融关键词"),
    (re.compile(r"(?i)结果|实验|指标|性能|评估"), "rows-2", "命中结果/实验关键词"),
    (re.compile(r"(?i)方法|流程|模型|架构|pipeline"), "cols-2-64", "命中方法/流程关键词"),
]


def _fail(msg: str) -> None:
    print(f"[extract_content] {msg}", file=sys.stderr)
    sys.exit(1)


def _normalize_text_item(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def _extract_image_hints_from_lines(lines: List[str]) -> List[str]:
    hints = set()
    for text in lines:
        content = _normalize_text_item(text)
        if not content:
            continue
        for img_path in MARKDOWN_IMAGE_RE.findall(content):
            hints.add(img_path)
        for img_path in IMAGE_PATH_RE.findall(content):
            hints.add(img_path)
        if FIGURE_MARK_RE.search(content):
            hints.add(content[:120])
    return sorted(hints)


def _build_layout_hints(sections: List[Dict[str, Any]]) -> List[Dict[str, str]]:
    layout_hints: List[Dict[str, str]] = []
    for sec in sections:
        title = sec.get("title", "未命名章节")
        texts = [title] + sec.get("bullets", []) + sec.get("paragraphs", [])
        merged = " ".join(_normalize_text_item(x) for x in texts if x)

        layout = "cols-2"
        reason = "默认文本页布局"
        for pattern, target_layout, target_reason in LAYOUT_RULES:
            if pattern.search(merged):
                layout = target_layout
                reason = target_reason
                break

        layout_hints.append({
            "section": title,
            "layout": layout,
            "reason": reason,
        })
    return layout_hints


def _read_md(path: str) -> Dict[str, Any]:
    with open(path, "r", encoding="utf-8") as f:
        text = f.read()
    return _read_md_text(text, os.path.splitext(os.path.basename(path))[0])


def _read_md_text(text: str, default_title: str) -> Dict[str, Any]:
    title = None
    sections = []
    current = None

    for line in text.splitlines():
        line = line.rstrip()
        if not line:
            continue

        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            heading = m.group(2).strip()
            if title is None:
                title = heading
            current = {"title": heading, "bullets": [], "paragraphs": []}
            sections.append(current)
            continue

        if line.lstrip().startswith(('-', '*')):
            bullet = line.lstrip()[1:].strip()
            if current is None:
                current = {"title": "内容", "bullets": [], "paragraphs": []}
                sections.append(current)
            current["bullets"].append(bullet)
        else:
            if current is None:
                current = {"title": "内容", "bullets": [], "paragraphs": []}
                sections.append(current)
            current["paragraphs"].append(line)

    image_hints = _extract_image_hints_from_lines(text.splitlines())

    return {
        "title": title or default_title,
        "sections": sections,
        "image_hints": image_hints,
    }


def _resolve_bundle_path(base: Path, paper_dir: Path, value: str) -> Path:
    p = Path(value)
    if p.is_absolute():
        return p
    candidate = (base / p)
    if candidate.exists():
        return candidate
    return (paper_dir / p)


def _read_arxiv_bundle(path: str) -> Dict[str, Any]:
    bundle_path = Path(path).resolve()
    with open(bundle_path, "r", encoding="utf-8") as f:
        bundle = json.load(f)

    paper_id = str(bundle.get("paper_id") or bundle_path.stem)
    paper_dir = Path(bundle.get("paper_dir") or bundle_path.parent).resolve()
    clean_text_value = bundle.get("clean_text_path")
    clean_text_path = None
    clean_text = ""
    if clean_text_value:
        clean_text_path = _resolve_bundle_path(bundle_path.parent, paper_dir, str(clean_text_value))
        if clean_text_path.exists():
            clean_text = clean_text_path.read_text(encoding="utf-8", errors="ignore")

    if not clean_text:
        merged_value = bundle.get("merged_tex_path")
        if merged_value:
            merged_path = _resolve_bundle_path(bundle_path.parent, paper_dir, str(merged_value))
            if merged_path.exists():
                clean_text = merged_path.read_text(encoding="utf-8", errors="ignore")

    md_like = _read_md_text(clean_text, paper_id)

    figures = bundle.get("figures", [])
    if not isinstance(figures, list):
        figures = []

    image_hints = list(md_like.get("image_hints", []))
    for fig in figures:
        if not isinstance(fig, dict):
            continue
        fig_path = fig.get("path")
        if not fig_path:
            continue
        hint_path = _resolve_bundle_path(bundle_path.parent, paper_dir, str(fig_path))
        image_hints.append(str(hint_path))

    source_url = bundle.get("source_url", "")
    if source_url:
        md_like.setdefault("citations", [])
        md_like["citations"].append(source_url)

    md_like["title"] = bundle.get("paper_id", md_like.get("title", paper_id))
    md_like["image_hints"] = sorted(set(image_hints))
    return md_like


def _read_pdf(path: str) -> Dict[str, Any]:
    try:
        from pypdf import PdfReader  # type: ignore
    except Exception:
        _fail("缺少依赖 pypdf。请先安装：pip install pypdf")

    reader = PdfReader(path)
    lines: List[str] = []
    for page in reader.pages:
        text = page.extract_text() or ""
        for ln in text.splitlines():
            ln = ln.strip()
            if ln:
                lines.append(ln)

    title = lines[0] if lines else os.path.splitext(os.path.basename(path))[0]

    sections = []
    current = {"title": "内容", "bullets": [], "paragraphs": []}
    sections.append(current)

    heading_re = re.compile(r"^\d+(\.\d+)*\s+|.+[：:]$")

    for ln in lines[1:]:
        if heading_re.match(ln):
            current = {"title": ln.rstrip('：:'), "bullets": [], "paragraphs": []}
            sections.append(current)
            continue
        if len(ln) <= 80:
            current["bullets"].append(ln)
        else:
            current["paragraphs"].append(ln)

    image_hints = _extract_image_hints_from_lines(lines)

    return {
        "title": title,
        "sections": sections,
        "image_hints": image_hints,
    }


def _read_pptx(path: str) -> Dict[str, Any]:
    try:
        from pptx import Presentation  # type: ignore
    except Exception:
        _fail("缺少依赖 python-pptx。请先安装：pip install python-pptx")

    prs = Presentation(path)
    sections = []
    title = None

    image_hints: List[str] = []
    all_text_lines: List[str] = []

    for idx, slide in enumerate(prs.slides, start=1):
        slide_title = None
        bullets = []
        paragraphs = []
        has_picture = False
        for shape in slide.shapes:
            if str(getattr(shape, "shape_type", "")) == "PICTURE (13)" or getattr(shape, "shape_type", None) == 13:
                has_picture = True
            if not hasattr(shape, "text"):
                continue
            text = shape.text.strip()
            if not text:
                continue
            all_text_lines.extend(text.splitlines())
            if slide_title is None and shape.has_text_frame:
                slide_title = text.splitlines()[0].strip()
                continue
            for ln in text.splitlines():
                ln = ln.strip()
                if not ln:
                    continue
                if ln.startswith(('-', '*')):
                    bullets.append(ln[1:].strip())
                else:
                    bullets.append(ln)
                if FIGURE_MARK_RE.search(ln):
                    image_hints.append(ln[:120])

        if title is None and slide_title:
            title = slide_title

        if has_picture:
            image_hints.append(f"Slide {idx} contains picture")

        sections.append({
            "title": slide_title or f"Slide {idx}",
            "bullets": bullets,
            "paragraphs": paragraphs,
        })

    image_hints.extend(_extract_image_hints_from_lines(all_text_lines))

    return {
        "title": title or os.path.splitext(os.path.basename(path))[0],
        "sections": sections,
        "image_hints": sorted(set(image_hints)),
    }


def _extract_citations(sections: List[Dict[str, Any]]) -> List[str]:
    citations = set()
    url_re = re.compile(r"https?://\S+")
    md_link_re = re.compile(r"\[[^\]]+\]\((https?://[^\)]+)\)")

    for sec in sections:
        for text in sec.get("bullets", []) + sec.get("paragraphs", []):
            for m in url_re.findall(text):
                citations.add(m)
            for m in md_link_re.findall(text):
                citations.add(m)

    return sorted(citations)


def main() -> None:
    parser = argparse.ArgumentParser(description="Extract content from PDF/PPTX/MD/arXiv-bundle JSON into JSON")
    parser.add_argument("input", help="input file path")
    parser.add_argument("--out", required=True, help="output json path")
    args = parser.parse_args()

    path = args.input
    if not os.path.exists(path):
        _fail(f"文件不存在: {path}")

    ext = os.path.splitext(path)[1].lower()
    if ext == ".md":
        data = _read_md(path)
    elif ext == ".pdf":
        data = _read_pdf(path)
    elif ext == ".pptx":
        data = _read_pptx(path)
    elif ext == ".json":
        data = _read_arxiv_bundle(path)
    else:
        _fail("不支持的文件类型，仅支持 .pdf / .pptx / .md / .json")

    data["source"] = os.path.basename(path)
    data["citations"] = _extract_citations(data.get("sections", []))
    data["image_hints"] = sorted(set(data.get("image_hints", [])))
    data["layout_hints"] = _build_layout_hints(data.get("sections", []))
    data.setdefault("figures", [])
    data.setdefault("tables", [])

    with open(args.out, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)

    print(f"[extract_content] 已写入: {args.out}")


if __name__ == "__main__":
    main()
